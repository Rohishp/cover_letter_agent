"""Spec-derived tests for T03 / SPEC.md Phase 1 check 5: cv.docx renders to <= 2 pages, no placeholders.

Page counts use a stand-in for cv.checks.convert_to_pdf (a real PDF with N blank pages), so they run
anywhere and fast. One smoke test uses the real converter and skips when none is available.
"""
import os

import pymupdf
import pytest
from docx import Document

from cv.checks import convert_to_pdf
from helpers import make_folder, verify_application as verify_without_pages
from verify_application import verify_application


def build_docx(path, lines=()):
    doc = Document()

    for line in lines:
        doc.add_paragraph(line)

    doc.save(str(path))


def fake_converter(monkeypatch, pages=None, garbage=False):
    """Replace the docx->PDF step: N blank pages, a corrupt PDF, or no converter (pages=None)."""

    def convert(docx_path):
        if pages is None and not garbage:
            return None, "no converter here"

        pdf = docx_path.with_suffix(".pdf")

        if garbage:
            pdf.write_bytes(b"not a pdf")
        else:
            document = pymupdf.open()

            for _ in range(pages):
                document.new_page()

            document.save(str(pdf))
            document.close()

        return pdf, "fake"

    monkeypatch.setattr("verify_application.convert_to_pdf", convert)


def cv_failures(folder, facts, pages=True):
    check = verify_application if pages else verify_without_pages

    return [f for f in check(folder, facts) if "cv.docx" in f]


@pytest.mark.parametrize("pages", [1, 2])
def test_cv_within_two_pages_is_clean(tmp_path, facts, monkeypatch, pages):
    fake_converter(monkeypatch, pages)

    assert cv_failures(make_folder(tmp_path), facts) == []


@pytest.mark.parametrize("pages", [3, 4])
def test_cv_over_two_pages_fails(tmp_path, facts, monkeypatch, pages):
    fake_converter(monkeypatch, pages)

    assert cv_failures(make_folder(tmp_path), facts) == [
        f"cv.docx: renders to {pages} pages, expected at most 2"
    ]


def test_no_converter_fails_closed(tmp_path, facts, monkeypatch):
    fake_converter(monkeypatch)

    assert cv_failures(make_folder(tmp_path), facts) == ["cv.docx: page count not checked, no converter here"]


def test_corrupt_render_fails_closed(tmp_path, facts, monkeypatch):
    fake_converter(monkeypatch, garbage=True)
    found = cv_failures(make_folder(tmp_path), facts)

    assert len(found) == 1 and "render failed" in found[0]


def test_render_leaves_no_pdf_in_the_folder(tmp_path, facts, monkeypatch):
    fake_converter(monkeypatch, 1)
    folder = make_folder(tmp_path)
    before = sorted(os.listdir(folder))
    verify_application(folder, facts)

    assert sorted(os.listdir(folder)) == before


@pytest.mark.parametrize("bad", ["[Company name]", "{{role}}", "TODO", "<company>"])
def test_placeholder_in_cv_fails(tmp_path, facts, bad):
    folder = make_folder(tmp_path)
    build_docx(folder / "cv.docx", [f"Applied to {bad} recently."])

    assert any("placeholder" in f for f in cv_failures(folder, facts, pages=False))


def test_placeholder_in_table_nested_table_header_and_footer_fails(tmp_path, facts):
    for where in ("table", "nested", "header", "footer"):
        folder = make_folder(tmp_path / where)
        doc = Document()

        if where == "table":
            doc.add_table(1, 1).cell(0, 0).text = "[x]"
        elif where == "nested":
            doc.add_table(1, 1).cell(0, 0).add_table(1, 1).cell(0, 0).text = "[x]"
        elif where == "header":
            doc.sections[0].header.paragraphs[0].text = "[x]"
        else:
            doc.sections[0].footer.paragraphs[0].text = "[x]"

        doc.save(str(folder / "cv.docx"))

        assert any("placeholder" in f for f in cv_failures(folder, facts, pages=False)), where


def test_placeholder_in_merged_cell_is_reported_once(tmp_path, facts):
    folder = make_folder(tmp_path)
    doc = Document()
    cells = doc.add_table(1, 3).rows[0].cells
    cells[0].merge(cells[2]).text = "[x]"
    doc.save(str(folder / "cv.docx"))

    assert len(cv_failures(folder, facts, pages=False)) == 1


def test_good_cv_has_no_failures_without_render(tmp_path, facts):
    folder = make_folder(tmp_path)
    build_docx(folder / "cv.docx", ["Built things at Acme Corp.", "Reach me at <me@example.com>."])

    assert cv_failures(folder, facts, pages=False) == []


def test_unreadable_cv_is_reported(tmp_path, facts):
    folder = make_folder(tmp_path)
    (folder / "cv.docx").write_bytes(b"not a docx at all")

    assert cv_failures(folder, facts, pages=False) == ["cv.docx is unreadable"]


def test_symlinked_cv_is_rejected(tmp_path, facts):
    folder = make_folder(tmp_path)
    target = tmp_path / "elsewhere.docx"
    (folder / "cv.docx").replace(target)

    try:
        (folder / "cv.docx").symlink_to(target)
    except OSError:
        pytest.skip("cannot create symlinks here")

    assert "cv.docx is a symlink" in verify_without_pages(folder, facts)


def test_declined_folder_skips_cv_checks(tmp_path, facts):
    assert verify_application(make_folder(tmp_path, declined=True), facts) == []


def test_real_converter_smoke(tmp_path, facts):
    probe = tmp_path / "probe.docx"
    Document().save(str(probe))
    pdf, _detail = convert_to_pdf(probe)

    if pdf is None:
        pytest.skip("no PDF converter (MS Word/LibreOffice) available")

    folder = make_folder(tmp_path / "app")
    doc = Document()

    for _ in range(2):
        doc.add_page_break()

    doc.save(str(folder / "cv.docx"))

    assert any("pages, expected at most 2" in f for f in cv_failures(folder, facts))
