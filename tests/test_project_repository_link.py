# Item 2: a project's repository link should render as a plain-text line
# after the tools/topic line, and only when a real repository URL exists.

from docx import Document

from cv.facts import ProjectLinks
from cv.render import DEFAULT_SECTION_ORDER, render_cv
from cv.schema import CVContent, EducationEntry, ProjectEntry
from cv.selection_rules import initial_state, materialize
from cv.selector import CVSelection
from models.jd_schema import ParsedJD


def _content(projects: list[ProjectEntry]) -> CVContent:
    return CVContent(
        full_name="Test Person",
        target_title="Test Title",
        location="Testville",
        contact_details=["test@example.com"],
        contact_links=[],
        profile="I built things.",
        skill_groups=[],
        experience=[],
        education=[
            EducationEntry(
                position="Degree", organisation="University", city="City",
                dates="2015 - 2018", bullets=[],
            )
        ],
        projects=projects,
        certifications=[],
        languages=[],
        coursework=[],
    )


def test_project_with_repository_link_renders_it(tmp_path):
    project = ProjectEntry(
        position="Widget Analyzer",
        organisation="Personal project",
        city="",
        dates="2024",
        tools="Python",
        bullets=["Analyzes widgets.", "Processed 1000 widgets per second."],
        repository_link="https://github.com/testuser/widget-analyzer",
    )

    docx_path = render_cv(_content([project]), tmp_path / "cv.docx", section_order=DEFAULT_SECTION_ORDER)
    doc = Document(str(docx_path))
    paragraphs = [p.text for p in doc.paragraphs]

    assert any("Repository: https://github.com/testuser/widget-analyzer" in p for p in paragraphs)

    tools_index = next(i for i, p in enumerate(paragraphs) if p.startswith("Tools: Python"))
    link_index = next(i for i, p in enumerate(paragraphs) if p.startswith("Repository:"))
    assert link_index == tools_index + 1


def test_project_without_repository_link_renders_nothing_extra(tmp_path):
    project = ProjectEntry(
        position="Gadget Tracker",
        organisation="Personal project",
        city="",
        dates="2024",
        tools="Python",
        bullets=["Tracks gadgets.", "Tracked gadgets in real time."],
    )

    docx_path = render_cv(_content([project]), tmp_path / "cv2.docx", section_order=DEFAULT_SECTION_ORDER)
    doc = Document(str(docx_path))
    paragraphs = [p.text for p in doc.paragraphs]

    assert not any("Repository:" in p for p in paragraphs)


def _selection() -> CVSelection:
    return CVSelection(
        section_order_reason="test",
        section_order=["skills", "experience", "projects", "education"],
        bullet_ranks=[],
        project_ranks=[],
        skill_group_ranks=[],
        target_title="Test Title",
    )


def test_materialize_threads_repository_link_from_facts(facts):
    facts.projects.entries[0].links = ProjectLinks(repository="https://github.com/x/y")
    jd = ParsedJD()
    selection = _selection()

    state, _report = initial_state(facts, selection, jd)
    content = materialize(facts, jd, selection, state, "A short profile.")

    widget_analyzer = next(p for p in content.projects if p.position == "Widget Analyzer")
    assert widget_analyzer.repository_link == "https://github.com/x/y"


def test_materialize_leaves_repository_link_none_without_facts_link(facts):
    jd = ParsedJD()
    selection = _selection()

    state, _report = initial_state(facts, selection, jd)
    content = materialize(facts, jd, selection, state, "A short profile.")

    assert all(p.repository_link is None for p in content.projects)
